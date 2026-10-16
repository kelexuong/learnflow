from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
from lf_common import LearnFlowError, workspace_path
from lf_parse import page_numbers
from lf_progress import load_progress, save_progress
from lf_sources import catalog, import_many, import_source, read_blocks, render_source_page, search, verify


class SourceTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="learnflow-source-test-")
        self.root = Path(self.temp.name)
        self.workspace = workspace_path(self.root / "study")

    def tearDown(self):
        self.temp.cleanup()

    def text(self, name="book.md", body="# 第一章\n\n充分条件保证结论成立。\n\n# 第二章\n\n必要条件不可缺少。"):
        path = self.root / name
        path.write_text(body, encoding="utf-8")
        return path

    def pdf(self, name="book.pdf", blank_second=False):
        from reportlab.pdfgen import canvas
        path = self.root / name
        document = canvas.Canvas(str(path))
        document.drawString(70, 750, "Chapter one: a sufficient condition guarantees its conclusion.")
        document.showPage()
        if not blank_second:
            document.drawString(70, 750, "Chapter two: a necessary condition is required.")
        document.showPage()
        document.save()
        return path

    def test_source_order_chinese_search_and_immutable_original(self):
        path = self.text()
        original = path.read_bytes()
        saved = import_source(self.workspace, path, source_id="book", role="book")
        first = read_blocks(self.workspace, "book", limit=2)
        second = read_blocks(self.workspace, "book", start=first["next_block"])
        self.assertIn("充分条件", first["blocks"][1]["text"])
        self.assertIn("第二章", second["blocks"][0]["section"])
        self.assertTrue(any("保证结论" in i["text"] for i in search(self.workspace, "充分条件")["results"]))
        self.assertEqual(path.read_bytes(), original)
        self.assertEqual(first["blocks"][0]["version"], saved["version"])
        self.assertFalse(first["reading_progress_changed"])

    def test_changed_version_preserves_history_and_learning_state(self):
        path = self.text()
        old = import_source(self.workspace, path, source_id="book")
        saved = save_progress(self.workspace, "logic", {"reading": {"version": old["version"]}, "next_step": "继续第一章"})
        path.write_text("# 修订第一章\n\n新版说明充分条件与必要条件。", encoding="utf-8")
        self.assertEqual(verify(self.workspace, "book")[0]["status"], "SOURCE_CHANGED")
        with self.assertRaises(LearnFlowError):
            import_source(self.workspace, path, source_id="book")
        with self.assertRaises(LearnFlowError):
            read_blocks(self.workspace, "book")
        current = import_source(self.workspace, path, source_id="book", new_version=True)
        self.assertNotEqual(current["version"], old["version"])
        self.assertEqual(len(catalog(self.workspace, "book")[0]["versions"]), 2)
        old_text = read_blocks(self.workspace, "book", version=old["version"], allow_stale=True)
        self.assertEqual(old_text["status"], "SOURCE_CHANGED")
        self.assertIn("第一章", old_text["blocks"][0]["text"])
        self.assertEqual(load_progress(self.workspace, "logic")["state"], saved["state"])

    def test_duplicate_bytes_keep_separate_provenance(self):
        a, b = self.text("a.md"), self.text("b.md")
        one = import_source(self.workspace, a, source_id="a")
        two = import_source(self.workspace, b, source_id="b")
        self.assertNotEqual(one["version"], two["version"])
        self.assertEqual(read_blocks(self.workspace, "b")["blocks"][0]["source_id"], "b")
        self.assertEqual(import_source(self.workspace, a, source_id="a")["status"], "UNCHANGED")

    def test_pdf_page_scope_and_original_preview(self):
        path = self.pdf()
        saved = import_source(self.workspace, path, source_id="pdf", pages="2")
        read = read_blocks(self.workspace, "pdf")
        self.assertEqual(read["blocks"][0]["location"]["page"], 2)
        self.assertNotIn("Chapter one", "".join(b["text"] for b in read["blocks"]))
        self.assertTrue(any("仅导入" in w for w in saved["warnings"]))
        from PIL import Image
        with Image.open(render_source_page(self.workspace, "pdf", 2)["path"]) as image:
            self.assertGreater(image.width, 800)

    def test_unreadable_page_is_neither_taught_nor_indexed(self):
        path = self.pdf(blank_second=True)
        with self.assertRaises(LearnFlowError):
            import_source(self.workspace, path, source_id="scan")
        self.assertEqual(catalog(self.workspace), [])
        import_source(self.workspace, path, source_id="scan", allow_partial=True)
        blocks = read_blocks(self.workspace, "scan")["blocks"]
        self.assertFalse(blocks[-1]["readable"])
        self.assertEqual(blocks[-1]["location"]["page"], 2)
        self.assertTrue(all(r["readable"] for r in search(self.workspace, "condition")["results"]))

    def test_ocr_bridge_marks_origin_and_missing_engine(self):
        path = self.pdf(blank_second=True)
        with self.assertRaises(LearnFlowError):
            import_source(self.workspace, path, source_id="scan", ocr=True, tesseract=str(self.root / "missing.exe"))
        with patch("lf_parse.ocr_page", return_value="Recognized scanned lesson text for review."):
            saved = import_source(self.workspace, path, source_id="scan", ocr=True)
        self.assertTrue(read_blocks(self.workspace, "scan")["blocks"][-1]["location"]["ocr"])
        self.assertTrue(any("OCR" in w for w in saved["warnings"]))

    def test_docx_table_picture_order_and_paragraph_locations(self):
        from docx import Document
        from docx.shared import Cm
        from PIL import Image
        document = Document()
        document.add_heading("第一节", level=1)
        document.add_paragraph("实际课堂案例与条件说明。")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text, table.cell(0, 1).text = "条件", "结论"
        image = self.root / "figure.png"
        Image.new("RGB", (40, 20), "green").save(image)
        document.add_picture(str(image), width=Cm(2))
        document.add_paragraph("最后一个段落。")
        path = self.root / "lesson.docx"
        document.save(path)
        saved = import_source(self.workspace, path, source_id="lesson", role="slides")
        blocks = read_blocks(self.workspace, "lesson")["blocks"]
        joined = "\n".join(b["text"] for b in blocks)
        self.assertLess(joined.index("条件 | 结论"), joined.index("最后一个段落"))
        self.assertTrue(any(b["location"].get("media") for b in blocks))
        self.assertTrue(list((Path(saved["snapshot_directory"]) / "media").iterdir()))
        self.assertTrue(all("page" not in b["location"] for b in blocks))

    def test_batch_continues_and_does_not_reimport_generated_outputs(self):
        inputs = self.workspace / "books"
        inputs.mkdir()
        (inputs / "good.md").write_text("# 教材\n\n真实教材正文。", encoding="utf-8")
        (inputs / "bad.txt").write_bytes(b"\xff\xff\xff")
        generated = self.workspace / "outputs"
        generated.mkdir()
        (generated / "generated.md").write_text("不应重新导入。", encoding="utf-8")
        result = import_many(self.workspace, [self.workspace])
        self.assertEqual((result["files"], result["errors"]), (2, 1))
        self.assertEqual(len(catalog(self.workspace)), 1)

    def test_bounded_read_can_resume_without_omissions(self):
        path = self.text(body="# 章节\n\n" + "\n\n".join(f"段落 {i}：" + "知识内容" * 200 for i in range(30)))
        import_source(self.workspace, path, source_id="long")
        seen, anchor = [], None
        while True:
            batch = read_blocks(self.workspace, "long", start=anchor, limit=3, max_chars=3000)
            seen.extend(b["block_id"] for b in batch["blocks"])
            self.assertLessEqual(batch["returned_chars"], 3000)
            anchor = batch["next_block"]
            if not anchor:
                break
        self.assertEqual((len(seen), len(set(seen))), (31, 31))

    def test_bad_ids_ranges_and_damaged_snapshot(self):
        with self.assertRaises(LearnFlowError):
            import_source(self.workspace, self.text(), source_id="../outside")
        for spec in ["0", "2-1", "1-4", "oops"]:
            with self.assertRaises(LearnFlowError):
                page_numbers(spec, 3)
        saved = import_source(self.workspace, self.text(), source_id="book")
        (Path(saved["snapshot_directory"]) / "blocks.jsonl").write_text("corrupted", encoding="utf-8")
        self.assertFalse(verify(self.workspace, "book")[0]["snapshot_ok"])

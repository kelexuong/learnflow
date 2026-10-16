from pathlib import Path
import json
import sys
import tempfile
import unittest
import zipfile

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
from lf_common import LearnFlowError, SKILL_ROOT, workspace_path
from lf_outputs import build_product, export_docx, export_visual


class OutputTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="learnflow-output-test-")
        self.workspace = workspace_path(Path(self.temp.name) / "study")

    def tearDown(self):
        self.temp.cleanup()

    def test_all_five_visual_types_have_real_embedded_data(self):
        examples = json.loads((SKILL_ROOT / "examples/visual-examples.json").read_text(encoding="utf-8"))["examples"]
        for data in examples:
            result = export_visual(self.workspace, data, f"visuals/{data['kind']}.html")
            text = Path(result["path"]).read_text(encoding="utf-8")
            self.assertNotIn("__DATA__", text)
            self.assertNotIn("__TITLE__", text)
            self.assertIn(data["title"], text)
            self.assertFalse(result["mastery_updated"])

    def test_visual_payload_cannot_escape_json_script(self):
        attack = '</script><script>window.PWNED=1</script>'
        data = {"kind": "flashcards", "title": attack, "cards": [{"front": attack, "back": "答案"}]}
        output = Path(export_visual(self.workspace, data, "visuals/safe.html")["path"]).read_text(encoding="utf-8")
        self.assertNotIn(attack, output)
        self.assertIn("\\u003c/script\\u003e", output)
        self.assertEqual(output.count("<script"), 2)

    def test_docx_exports_actual_headings_lists_code_and_table(self):
        markdown = self.workspace / "notes.md"
        original = "# 条件关系学习笔记\n\n本次围绕实际资料区分两类条件。\n\n## 核心区别\n\n- **充分条件**：保证结论。\n- 必要条件：不可缺少。\n\n| 关系 | 含义 |\n|---|---|\n| P→Q | P 是 Q 的充分条件 |\n\n```text\nP -> Q\n```\n"
        markdown.write_text(original, encoding="utf-8")
        saved = export_docx(self.workspace, markdown, "outputs/notes.docx")
        from docx import Document
        document = Document(saved["path"])
        self.assertEqual(document.paragraphs[0].style.name, "Title")
        self.assertIn("条件关系", document.paragraphs[0].text)
        self.assertEqual(document.tables[0].cell(1, 0).text, "P→Q")
        self.assertTrue(any(p.text == "P -> Q" for p in document.paragraphs))
        self.assertEqual(markdown.read_text(encoding="utf-8"), original)

    def test_output_cannot_overwrite_or_escape_workspace(self):
        data = {"kind": "flashcards", "title": "题目", "cards": [{"front": "问", "back": "答"}]}
        export_visual(self.workspace, data, "once.html")
        with self.assertRaises(LearnFlowError):
            export_visual(self.workspace, data, "once.html")
        with self.assertRaises(LearnFlowError):
            export_visual(self.workspace, data, "../outside.html")
        with self.assertRaises(LearnFlowError):
            export_visual(self.workspace, data, "wrong.py")

    def test_packaging_is_explicit_and_excludes_unlisted_files(self):
        folder = self.workspace / "products/ready"
        folder.mkdir(parents=True)
        (folder / "notes.md").write_text("# 原创总结\n\n这是准备复用的学习结果。", encoding="utf-8")
        (folder / "private.md").write_text("未列入清单。", encoding="utf-8")
        manifest = {"kind": "knowledge-pack", "name": "study-notes", "files": [{"path": "products/ready/notes.md", "as": "notes.md"}]}
        result = build_product(self.workspace, manifest, "products/notes.zip")
        with zipfile.ZipFile(result["path"]) as archive:
            self.assertEqual(set(archive.namelist()), {"study-notes/notes.md", "study-notes/product-manifest.json"})
        self.assertFalse(result["published"])

    def test_packaging_rejects_private_sources_even_when_renamed(self):
        source = self.workspace / "sources/private.md"
        source.parent.mkdir()
        source.write_text("原资料", encoding="utf-8")
        for destination in ["notes.md", "../escape.md", "C:/escape.md"]:
            manifest = {"kind": "knowledge-pack", "name": "study", "files": [{"path": "sources/private.md", "as": destination}]}
            with self.assertRaises(LearnFlowError):
                build_product(self.workspace, manifest, "products/denied.zip")
        self.assertFalse((self.workspace / "products/denied.zip").exists())

    def test_packaging_rejects_reserved_and_case_colliding_names(self):
        source = self.workspace / "notes.md"
        source.write_text("原创总结", encoding="utf-8")
        for rows in [
            [{"path": "notes.md", "as": "product-manifest.json"}],
            [{"path": "notes.md", "as": "Notes.md"}, {"path": "notes.md", "as": "notes.md"}],
        ]:
            manifest = {"kind": "knowledge-pack", "name": "notes", "files": rows}
            with self.assertRaises(LearnFlowError):
                build_product(self.workspace, manifest, "denied.zip")
        self.assertFalse((self.workspace / "denied.zip").exists())

    def test_skill_product_requires_all_local_references(self):
        folder = self.workspace / "products/ready"
        folder.mkdir(parents=True)
        (folder / "SKILL.md").write_text('---\nname: lesson-tools\ndescription: Explain original lesson tools.\n---\n\n# Tools\n\nRead [guide](guide.md).', encoding="utf-8")
        manifest = {"kind": "skill", "name": "lesson-tools", "files": [{"path": "products/ready/SKILL.md", "as": "SKILL.md"}]}
        with self.assertRaises(LearnFlowError):
            build_product(self.workspace, manifest, "products/lesson.zip")
        (folder / "guide.md").write_text("# Guide\n\n实际可复用内容。", encoding="utf-8")
        manifest["files"].append({"path": "products/ready/guide.md", "as": "guide.md"})
        result = build_product(self.workspace, manifest, "products/lesson.zip")
        self.assertTrue(Path(result["path"]).is_file())

    def test_invalid_visual_graph_is_rejected_before_writing(self):
        data = {"kind": "concept-map", "title": "错误关系", "nodes": [{"id": "a", "title": "A", "body": "内容"}], "edges": [{"from": "a", "to": "missing"}]}
        with self.assertRaises(LearnFlowError):
            export_visual(self.workspace, data, "visuals/invalid.html")
        self.assertFalse((self.workspace / "visuals/invalid.html").exists())

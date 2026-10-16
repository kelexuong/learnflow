"""Optional learning outputs; no teaching policy or evidence is inferred here."""

from __future__ import annotations

import hashlib
import html
import json
from pathlib import Path
import re
import zipfile

from lf_common import LearnFlowError, SKILL_ROOT, atomic_json, digest_file


def output_path(workspace: Path, value: str | Path) -> Path:
    path = Path(value).expanduser()
    path = (workspace / path).resolve() if not path.is_absolute() else path.resolve()
    if not path.is_relative_to(workspace) or path == SKILL_ROOT or path.is_relative_to(SKILL_ROOT):
        raise LearnFlowError("学习产物请保存在 --workspace 内，并位于 Skill 目录外。")
    if path.exists():
        raise LearnFlowError("输出路径已经存在，请选择新路径以保留原文件。")
    path.parent.mkdir(parents=True, exist_ok=True)
    return path


def string(value, label):
    if not isinstance(value, str) or not value.strip():
        raise LearnFlowError(f"{label} 需要非空文字。")
    return value


def validate_visual(data: dict):
    if not isinstance(data, dict):
        raise LearnFlowError("交互页数据须为 JSON 对象。")
    string(data.get("title"), "title")
    kind = data.get("kind")
    key = {"concept-map": "nodes", "process": "steps", "compare": "items",
           "flashcards": "cards", "quiz": "questions"}.get(kind)
    if key is None:
        raise LearnFlowError("kind 须为 concept-map/process/compare/flashcards/quiz。")
    rows = data.get(key)
    if not isinstance(rows, list) or not 1 <= len(rows) <= 60:
        raise LearnFlowError(f"{key} 须包含 1–60 项；较大的内容请按学习目标分开。")
    for row in rows:
        if not isinstance(row, dict):
            raise LearnFlowError("交互数据项须为对象。")
        fields = ("front", "back") if kind == "flashcards" else (
            ("prompt", "answer") if kind == "quiz" else ("title", "body"))
        for field in fields:
            string(row.get(field), field)
        if "source_refs" in row and (not isinstance(row["source_refs"], list) or
                                     any(not isinstance(r, str) for r in row["source_refs"])):
            raise LearnFlowError("source_refs 须为来源定位文字列表。")
    if kind == "concept-map":
        ids = [string(row.get("id"), "node.id") for row in rows]
        if len(ids) != len(set(ids)):
            raise LearnFlowError("知识节点 ID 不可重复。")
        for edge in data.get("edges", []):
            if not isinstance(edge, dict) or edge.get("from") not in ids or edge.get("to") not in ids:
                raise LearnFlowError("关系边引用了未知节点。")


def export_visual(workspace: Path, data: dict, output: str | Path) -> dict:
    validate_visual(data)
    if Path(output).suffix.lower() != ".html":
        raise LearnFlowError("交互学习页输出须为 .html。")
    path = output_path(workspace, output)
    payload = json.dumps(data, ensure_ascii=False, allow_nan=False)
    payload = payload.replace("&", "\\u0026").replace("<", "\\u003c").replace(">", "\\u003e").replace("\u2028", "\\u2028").replace("\u2029", "\\u2029")
    template = (SKILL_ROOT / "assets" / "learning-page.html").read_text(encoding="utf-8")
    content = template.replace("__TITLE__", html.escape(data["title"], quote=True)).replace("__DATA__", payload)
    path.write_text(content, encoding="utf-8", newline="\n")
    return {"path": str(path), "kind": data["kind"], "offline": True,
            "sha256": digest_file(path), "mastery_updated": False}


def inline_runs(paragraph, text: str):
    for part in re.split(r"(\*\*.+?\*\*|`[^`]+`)", text):
        run = paragraph.add_run(part[2:-2] if part.startswith("**") and part.endswith("**") else
                                part[1:-1] if part.startswith("`") and part.endswith("`") else part)
        if part.startswith("**") and part.endswith("**"):
            run.bold = True
        elif part.startswith("`") and part.endswith("`"):
            run.font.name = "Consolas"


def export_docx(workspace: Path, markdown: Path, output: str | Path) -> dict:
    try:
        from docx import Document
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.shared import Cm, Pt, RGBColor
    except ImportError as exc:
        raise LearnFlowError("DOCX 导出需要 python-docx；请安装 requirements.txt。") from exc
    markdown = markdown.expanduser().resolve()
    if not markdown.is_file():
        raise LearnFlowError("未找到准备导出的 Markdown 文件。")
    if Path(output).suffix.lower() != ".docx":
        raise LearnFlowError("Word 输出须为 .docx。")
    lines = markdown.read_text(encoding="utf-8-sig").splitlines()
    if not any(line.strip() for line in lines):
        raise LearnFlowError("Markdown 没有正文，未创建空白手册。")
    path = output_path(workspace, output)
    document = Document()
    section = document.sections[0]
    section.top_margin = section.bottom_margin = Cm(2)
    section.left_margin = section.right_margin = Cm(2.2)
    for name in ["Normal", "Title", "Heading 1", "Heading 2", "Heading 3"]:
        style = document.styles[name]
        style.font.name = "Calibri"
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    document.styles["Normal"].font.size = Pt(11)
    document.styles["Normal"].paragraph_format.space_after = Pt(6)
    document.styles["Normal"].paragraph_format.line_spacing = 1.15
    document.styles["Title"].font.size = Pt(24)
    code, first_heading, index = False, True, 0
    while index < len(lines):
        line = lines[index]
        if line.strip().startswith("```") or line.strip().startswith("~~~"):
            code = not code
            index += 1
            continue
        if code:
            paragraph = document.add_paragraph(style="No Spacing")
            run = paragraph.add_run(line)
            run.font.name, run.font.size = "Consolas", Pt(9)
            index += 1
            continue
        if "|" in line and index + 1 < len(lines) and re.fullmatch(r"\s*\|?\s*:?-+:?\s*(?:\|\s*:?-+:?\s*)+\|?\s*", lines[index + 1]):
            def cells(row):
                return [part.strip() for part in row.strip().strip("|").split("|")]
            header = cells(line)
            rows = [header]
            index += 2
            while index < len(lines) and "|" in lines[index] and lines[index].strip():
                row = cells(lines[index])
                if len(row) != len(header):
                    raise LearnFlowError("Markdown 表格列数不一致，请先修正。")
                rows.append(row)
                index += 1
            table = document.add_table(rows=len(rows), cols=len(header))
            table.style = "Table Grid"
            for r, values in enumerate(rows):
                for c, text in enumerate(values):
                    paragraph = table.cell(r, c).paragraphs[0]
                    inline_runs(paragraph, text)
                    if r == 0:
                        for run in paragraph.runs:
                            run.bold = True
            repeat = OxmlElement("w:tblHeader")
            table.rows[0]._tr.get_or_add_trPr().append(repeat)
            continue
        heading = re.match(r"^(#{1,6})\s+(.+)$", line)
        if heading:
            level = len(heading[1])
            if first_heading and level == 1:
                paragraph = document.add_paragraph(style="Title")
            else:
                paragraph = document.add_paragraph(style=f"Heading {min(level, 3)}")
            inline_runs(paragraph, heading[2])
            first_heading = False
        elif line.strip():
            bullet = re.match(r"^\s*[-*+]\s+(.+)$", line)
            ordered = re.match(r"^\s*\d+[.)]\s+(.+)$", line)
            paragraph = document.add_paragraph(style="List Bullet" if bullet else "List Number" if ordered else "Normal")
            inline_runs(paragraph, bullet[1] if bullet else ordered[1] if ordered else line)
        index += 1
    document.core_properties.title = next((line[2:].strip() for line in lines if line.startswith("# ")), markdown.stem)
    document.save(path)
    # Reopen the package to catch malformed output, beyond a mere existence check.
    if not Document(path).paragraphs:
        raise LearnFlowError("导出的 DOCX 没有正文。")
    return {"path": str(path), "sha256": digest_file(path),
            "note": "已保留文字、标题、列表和简单表格；数学布局、复杂图文请对照原文人工核对。"}


def build_product(workspace: Path, manifest: dict, output: str | Path) -> dict:
    """Package only explicitly supplied public-ready files, never a whole study folder."""
    if not isinstance(manifest, dict):
        raise LearnFlowError("成果清单须为对象。")
    kind = manifest.get("kind")
    if kind not in {"knowledge-pack", "skill"}:
        raise LearnFlowError("成果类型须为 knowledge-pack 或 skill。")
    name = manifest.get("name", "")
    if not re.fullmatch(r"[a-z0-9]+(?:-[a-z0-9]+)*", name) or len(name) > 64:
        raise LearnFlowError("成果名须为小写英文与连字符，最多 64 字符。")
    rows = manifest.get("files")
    if not isinstance(rows, list) or not rows:
        raise LearnFlowError("须显式列出准备公开复用的文件；不会自动打包学习档案。")
    selected = []
    names = set()
    for item in rows:
        if not isinstance(item, dict) or not isinstance(item.get("path"), str) or not isinstance(item.get("as"), str):
            raise LearnFlowError("每个文件需要 path 和 as。")
        source = (workspace / item["path"]).resolve()
        target = Path(item["as"])
        if not source.is_relative_to(workspace) or not source.is_file() or source.is_symlink():
            raise LearnFlowError("成果输入须为资料工作区内明确列出的普通文件。")
        relative_source = source.relative_to(workspace)
        if any(part in {"sources", "progress", "records", "learning-records"} for part in relative_source.parts) or source.name.startswith(".env") or source.suffix in {".sqlite3", ".db"}:
            raise LearnFlowError("请先准备可公开的成果文件，不直接打包资料快照、学习档案或环境配置。")
        if target.is_absolute() or ".." in target.parts or not target.parts or ":" in item["as"] or "\\" in item["as"]:
            raise LearnFlowError("成果内部路径须为安全的相对路径。")
        if target.parts[0] in {"sources", "progress", "records", "learning-records"} or target.name in {"catalog.sqlite3", ".env"}:
            raise LearnFlowError("学习档案、原资料和环境配置不能作为默认可公开成果。")
        normalized = target.as_posix()
        if normalized.casefold() == "product-manifest.json":
            raise LearnFlowError("product-manifest.json 由打包工具生成，请为成果正文选择其他名称。")
        if normalized.casefold() in {n.casefold() for n in names}:
            raise LearnFlowError("成果内部文件名重复。")
        names.add(normalized)
        selected.append((source, normalized))
    if kind == "skill":
        if "SKILL.md" not in names:
            raise LearnFlowError("Skill 成果必须包含实际编写完成的 SKILL.md。")
        body = next(p.read_text(encoding="utf-8-sig") for p, n in selected if n == "SKILL.md")
        front = re.match(r"\A---\r?\n(.*?)\r?\n---(?:\r?\n|$)", body, re.S)
        try:
            import yaml
            meta = yaml.safe_load(front[1]) if front else None
        except ImportError as exc:
            raise LearnFlowError("Skill 成果校验需要 PyYAML；请安装 requirements.txt。") from exc
        except yaml.YAMLError as exc:
            raise LearnFlowError("成果前置元数据不是有效 YAML。") from exc
        if not isinstance(meta, dict) or meta.get("name") != name or not isinstance(meta.get("description"), str) or not meta["description"].strip():
            raise LearnFlowError("Skill 成果须有与成果名一致的 name 和非空 description 前置元数据。")
        for source, target in selected:
            if target.endswith(".md"):
                for link in re.findall(r"!?\[[^\]\n]*\]\(([^)\s]+)\)", source.read_text(encoding="utf-8-sig")):
                    from urllib.parse import urlsplit, unquote
                    import posixpath
                    parts = urlsplit(link)
                    if parts.scheme or parts.netloc or not parts.path:
                        continue
                    resolved = posixpath.normpath(posixpath.join(posixpath.dirname(target), unquote(parts.path)))
                    if resolved.startswith("../") or resolved not in names:
                        raise LearnFlowError("Skill 成果有未打包的本地引用：" + link)
    if Path(output).suffix.lower() != ".zip":
        raise LearnFlowError("成果包输出须为 .zip。")
    path = output_path(workspace, output)
    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for source, target in selected:
            archive.write(source, name + "/" + target)
        archive.writestr(name + "/product-manifest.json", json.dumps(
            {"name": name, "kind": kind, "files": [{"path": t, "sha256": digest_file(p)} for p, t in selected]}, ensure_ascii=False, indent=2))
    return {"path": str(path), "sha256": digest_file(path), "files": len(selected), "published": False}
